"""Build the submission report from Markdown: PDF (via headless Chrome) and DOCX.

Usage: python scripts/build_report.py [--src docs/report.md] [--name report]
       python scripts/build_report.py --src docs/gate_defense_memo.md --name gate_defense_memo
Output: reports/<name>.html, reports/<name>.pdf, reports/<name>.docx
Needs: requirements-docs.txt (Markdown, python-docx) and Google Chrome / Chromium.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
from pathlib import Path

import markdown

CSS = """
@page { size: A4; margin: 15mm 15mm 14mm 15mm; }
body { font-family: 'Helvetica Neue', Arial, sans-serif; font-size: 9.6pt; line-height: 1.38; color: #1d2731; }
h1 { font-size: 16pt; margin: 0 0 2pt 0; color: #264653; }
h2 { font-size: 11.5pt; margin: 9pt 0 3pt 0; color: #264653; border-bottom: 1px solid #d5dde3; padding-bottom: 1pt; }
h3 { font-size: 10pt; margin: 6pt 0 2pt 0; color: #2a6f66; }
p { margin: 2pt 0 4pt 0; text-align: justify; }
ul, ol { margin: 2pt 0 4pt 0; padding-left: 15pt; }
li { margin: 0 0 1.5pt 0; }
table { border-collapse: collapse; width: 100%; margin: 3pt 0 5pt 0; font-size: 8.3pt; }
th { background: #264653; color: white; text-align: left; padding: 2pt 4pt; font-weight: 600; }
td { border-bottom: 1px solid #e3e8ec; padding: 1.6pt 4pt; vertical-align: top; }
tr:nth-child(even) td { background: #f6f8f9; }
img { max-width: 100%; display: block; margin: 3pt auto; }
img.half { max-width: 49%; display: inline-block; }
code { font-family: Menlo, monospace; font-size: 8pt; background: #f1f3f5; padding: 0 2pt; border-radius: 2px; }
.meta { color: #5c6b77; font-size: 8.6pt; margin-bottom: 5pt; }
blockquote { margin: 3pt 0; padding: 2pt 8pt; border-left: 3px solid #2a9d8f; background: #eef7f5; }
strong { color: #13232f; }
"""


def find_chrome() -> str:
    for c in ("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", "google-chrome", "chromium",
              "chromium-browser"):
        if Path(c).exists() or shutil.which(c):
            return c
    raise SystemExit("Chrome/Chromium not found: needed to print the PDF")


def to_html(md_text: str, base: Path, title: str) -> str:
    body = markdown.markdown(md_text, extensions=["tables", "fenced_code", "attr_list", "sane_lists"])
    # Make image paths absolute so Chrome can load them from the HTML location.
    body = re.sub(r'src="([^"]+)"', lambda m: f'src="{(base / m.group(1)).resolve().as_uri()}"', body)
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title><style>{CSS}</style></head>" \
           f"<body>{body}</body></html>"


def to_docx(md_text: str, base: Path, out: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin = s.bottom_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9.5)

    def add_runs(par, text: str) -> None:
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
            if not part:
                continue
            if part.startswith("**"):
                par.add_run(part[2:-2]).bold = True
            elif part.startswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                par.add_run(part[1:-1]).italic = True
            else:
                par.add_run(part)

    def clean(t: str) -> str:
        t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
        return t.replace("{: .half}", "").strip()

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i].strip()):
                    rows.append([clean(c) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row[: len(rows[0])]):
                    p = table.cell(r, c).paragraphs[0]
                    add_runs(p, cell)
                    for run in p.runs:
                        run.font.size = Pt(8)
                        if r == 0:
                            run.bold = True
            doc.add_paragraph()
            continue
        m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", line.strip())
        if m:
            img = (base / m.group(1)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Cm(17.5 if "{: .half}" not in line else 8.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        h = re.match(r"^(#{1,3}) (.*)", line)
        if h:
            par = doc.add_heading(clean(h.group(2)), level=len(h.group(1)))
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x26, 0x46, 0x53)
        elif re.match(r"^\s*[-*] ", line):
            add_runs(doc.add_paragraph(style="List Bullet"), clean(re.sub(r"^\s*[-*] ", "", line)))
        elif re.match(r"^\s*\d+\. ", line):
            add_runs(doc.add_paragraph(style="List Number"), clean(re.sub(r"^\s*\d+\. ", "", line)))
        elif line.startswith("<p class=\"meta\">"):
            p = doc.add_paragraph()
            add_runs(p, clean(re.sub(r"<[^>]+>", "", line)))
        elif line.strip() and not line.startswith("<"):
            add_runs(doc.add_paragraph(), clean(line.lstrip("> ")))
        i += 1
    doc.save(out)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default="docs/report.md")
    ap.add_argument("--name", default="report")
    args = ap.parse_args()
    src = Path(args.src)
    md_text = src.read_text()
    out_dir = Path("reports")
    html_path = out_dir / f"{args.name}.html"
    html_path.write_text(to_html(md_text, src.parent, args.name))
    pdf_path = out_dir / f"{args.name}.pdf"
    subprocess.run([find_chrome(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf_path.resolve()}", html_path.resolve().as_uri()],
                   check=True, capture_output=True)
    to_docx(md_text, src.parent, out_dir / f"{args.name}.docx")
    pages = len(re.findall(rb"/Type\s*/Page[^s]", pdf_path.read_bytes()))
    print(f"{pdf_path} ({pages} pages), {out_dir / f'{args.name}.docx'}, {html_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
